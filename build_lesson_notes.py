from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/alina/Downloads/Конспект — Отношения — сценарий.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(100, 108, 118)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK, 8, 4),
]:
    s = styles[name]
    s.font.name = "Calibri"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for name in ["List Bullet", "List Number"]:
    s = styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(11)
    s.paragraph_format.left_indent = Inches(0.5)
    s.paragraph_format.first_line_indent = Inches(-0.25)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.1

header = section.header.paragraphs[0]
header.text = "SILVERWIND  •  КОНСПЕКТ УРОКА"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in header.runs:
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Страница ")
run.font.size = Pt(9)
run.font.color.rgb = MUTED
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ОТНОШЕНИЯ — СЦЕНАРИЙ")
r.bold = True
r.font.name = "Calibri"
r.font.size = Pt(25)
r.font.color.rgb = DARK

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("Конспект урока и практика по переписыванию внутренних программ")
r.font.size = Pt(12)
r.font.color.rgb = MUTED

def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        a, b = text.split(bold_prefix, 1)
        p.add_run(bold_prefix).bold = True
        p.add_run(b)
    else:
        p.add_run(text)
    return p

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def numbered(text):
    doc.add_paragraph(text, style="List Number")

def quote(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.08
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    pPr.append(shd)
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = DARK

doc.add_heading("Главная идея", level=1)
para("В логике урока отношения — это не отдельная внешняя реальность, а сценарий, отражающий внутренние знания и представления о себе, других людях, любви, дружбе и близости.")
para("Внешние отношения сравниваются с фильмом на стене: фильм — уже следствие, а «проектором» выступает сознание человека. То, что проявлено сейчас, не является окончательной правдой: это результат внимания, направленного на определённые знания и сценарии.")
quote("Изменяя внутреннюю систему наблюдения, человек выбирает другой сценарий отношений.")

doc.add_heading("Диагностика сценария отношений", level=1)
para("Для практики нужно выбрать одну конкретную группу: мужчины, женщины, партнёры, подруги, клиенты, работодатели, родители или дети.")

doc.add_heading("1. Что для меня означают эти отношения?", level=2)
for x in ["Что такое отношения с этими людьми?", "Какими они обычно бывают?", "Что в них чаще всего происходит?", "Могут ли они быть долгими и спокойными?", "Являются ли отношения трудом и постоянным преодолением?", "Возможна ли близость без боли, конфликтов и предательства?"]:
    bullet(x)
para("Записывать нужно не социально правильные ответы, а то, во что человек действительно верит.")

doc.add_heading("2. Что я знаю об этой группе людей?", level=2)
for x in ["Какие они?", "Как они обычно поступают?", "Чего от них можно ожидать?", "Что они делают и чего, по моему знанию, никогда не делают?"]:
    bullet(x)
para("Чтобы обнаружить настоящее убеждение, полезно вспомнить, что вы недавно говорили об этих людях подруге или близкому человеку. Спонтанные разговоры часто точнее показывают внутренние программы.")

doc.add_heading("3. Кто я внутри этих отношений?", level=2)
for x in ["Какая я рядом с ними?", "Как они меня воспринимают?", "Какое место мне обычно достаётся?", "Чего я заслуживаю и чего не заслуживаю?", "Какой образ себя я поддерживаю в этом контакте?"]:
    bullet(x)

doc.add_heading("Отношение к себе как ядро", level=1)
para("Отдельная часть практики — наблюдение за тем, как человек обращается с собой. Важно отследить:")
for x in ["как вы разговариваете с собой после ошибок;", "поддерживаете или оскорбляете себя;", "ставите ли свои потребности на последнее место;", "разрешаете ли себе отдых, удовольствие и заботу;", "отдаёте ли себе только то, что осталось после других;", "из какого состояния взаимодействуете с людьми."]:
    bullet(x)
para("Наблюдать за собой желательно несколько дней или неделю, потому что отношение к себе часто проявляется автоматически.")
quote("Три направления наблюдения: как я отношусь к себе; как я отношусь к другим; как другие относятся ко мне.")

doc.add_heading("Программы и чужие знания", level=1)
para("Обнаруженные представления — не обязательно осознанно выбранная истина. Многие из них могли быть переданы семьёй, окружением или прошлым опытом.")
for x in ["Кто мне это сказал?", "Откуда этот человек это узнал?", "Почему я считаю это своим знанием?", "Выбирала ли я сама смотреть на отношения именно так?"]:
    bullet(x)
quote("Я вижу это убеждение, но оно не является мной. Я не выбирала его осознанно и могу перестать направлять в него внимание.")

doc.add_heading("Переписывание сценария", level=1)
para("Старые знания предлагается символически сложить в старый чемодан и отпустить. Новый чемодан заполняется осознанно выбранными наблюдениями.")
for x in ["В мире много достойных и заботливых мужчин.", "Существуют женщины, которые искренне хотят дружить со мной.", "Есть люди, которым подходит именно мой способ проявления.", "Есть клиенты, которым нужен именно мой продукт.", "Я могу создавать отношения, соответствующие выбранному мной знанию."]:
    bullet(x)
para("Новая программа получает силу не просто потому, что она звучит позитивно, а потому, что теперь это осознанный выбор самого человека.")

doc.add_heading("Основные программы плохого отношения к себе", level=1)
doc.add_heading("1. Условная ценность", level=2)
para("Человек обещает начать ценить и уважать себя после выполнения условия: когда станет богатым, реализуется, похудеет, вступит в отношения, создаст семью или получит признание.")
quote("Моя ценность существует до денег, отношений, внешности и результатов. Всё это может быть моим желанием, но не условием права любить и уважать себя.")

doc.add_heading("2. Самонаказание", level=2)
para("Когда условие ценности не выполнено, человек оскорбляет себя, лишает заботы, проявляет к себе насилие или считает, что пока не заслуживает хорошего отношения.")
quote("Результат не создаёт мою ценность. Опыт рождается из той точки, из которой я уже существую и выбираю.")

doc.add_heading("3. Самопожертвование", level=2)
para("В этой программе любовь приравнивается к жертве, нужно быть удобной ради признания, а чужие потребности и комфорт ставятся выше собственных.")
quote("Мне не нужно уменьшать себя, чтобы стать источником любви. Другой человек не является условием моей ценности.")

doc.add_heading("Иерархия «я и мой опыт»", level=1)
para("Деньги, отношения, тело, реализация и признание не могут быть больше человека: в философии урока они появляются в его опыте благодаря его вниманию и присутствию.")
para("Желать денег, отношений, красоты и реализации можно. Искажение начинается тогда, когда желание становится условием права уважать себя.")
quote("Не я завишу от опыта как от источника своей ценности — опыт зависит от моего присутствия и выбора.")

doc.add_heading("Практика после урока", level=1)
steps = [
    "Выберите одну группу людей или одну сферу отношений.",
    "Выпишите всё, что вы знаете об этих отношениях.",
    "Опишите собирательный образ этих людей.",
    "Определите, кем вы считаете себя рядом с ними.",
    "Несколько дней наблюдайте за своим внутренним диалогом.",
    "Найдите условия, после выполнения которых обещаете начать себя ценить.",
    "Для каждой деструктивной программы спросите: «Кто мне это передал и выбирала ли я это сама?»",
    "Разотождествитесь со старым знанием.",
    "Сформулируйте новое наблюдение, которое выбираете осознанно.",
    "Начните относиться к себе из новой программы сейчас, не ожидая внешнего подтверждения.",
]
for x in steps:
    numbered(x)

doc.add_heading("Ключевые формулы урока", level=1)
for x in [
    "Отношения — проявленный сценарий моих внутренних знаний.",
    "То, что я наблюдаю сейчас, не является зафиксированной реальностью.",
    "Переданное мне убеждение — не обязательно моё собственное знание.",
    "Я могу осознанно выбрать новую точку наблюдения.",
    "Моя ценность не создаётся опытом. Опыт создаётся из меня.",
    "Деньги, отношения, внешность и реализация — мои желания, а не условия права любить себя.",
    "Мне не нужно жертвовать собой, чтобы получить любовь.",
    "Я существую раньше любого результата — поэтому моя ценность безусловна.",
]:
    quote(x)

doc.core_properties.title = "Конспект урока «Отношения — сценарий»"
doc.core_properties.subject = "Отношения, внутренние программы и безусловная ценность"
doc.save(OUT)
print(OUT)
